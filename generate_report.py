from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

document = Document()

# Add a title
title = document.add_heading('LAPORAN KEGIATAN DAN PERTANGGUNGJAWABAN', 1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = document.add_paragraph('PEMBUATAN APLIKASI LAPOR FTI')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].bold = True

document.add_paragraph()

# Details
p = document.add_paragraph()
p.add_run('A. DATA PELAKSANA TUGAS').bold = True
document.add_paragraph('Berdasarkan Surat Tugas No. 300/Dek-FTI/60/FTI/VIII/2026 tanggal 12 Agustus 2026, dengan ini dilaporkan pelaksanaan tugas sebagai berikut:')
document.add_paragraph('1. Nama\t\t: Zulfahmi Nur Kesuma Atmaja, A.Md.')
document.add_paragraph('2. Tugas\t\t: Pembuatan Aplikasi Lapor FTI')
document.add_paragraph('3. Waktu\t\t: 7 Agustus 2026 s.d. 7 Desember 2026 (122 Hari)')

document.add_paragraph()
p = document.add_paragraph()
p.add_run('B. HASIL PELAKSANAAN TUGAS').bold = True
document.add_paragraph('Tugas pembuatan Aplikasi Lapor FTI (Fakultas Teknologi Industri Universitas Islam Indonesia) telah dilaksanakan pada rentang waktu yang telah ditentukan. Adapun beberapa capaian dan status pengembangan aplikasi saat ini adalah sebagai berikut:')
p_list = document.add_paragraph('- Pengembangan struktur basis data dan antarmuka (UI/UX) aplikasi.')
p_list.paragraph_format.left_indent = Inches(0.25)
p_list2 = document.add_paragraph('- Implementasi fitur utama pelaporan untuk memfasilitasi aduan/keluhan.')
p_list2.paragraph_format.left_indent = Inches(0.25)
p_list3 = document.add_paragraph('- Pengujian fungsionalitas dasar aplikasi.')
p_list3.paragraph_format.left_indent = Inches(0.25)

document.add_paragraph()
p = document.add_paragraph()
p.add_run('C. KENDALA DAN PEKERJAAN YANG BELUM SELESAI (BACKLOG)').bold = True
document.add_paragraph('Dalam proses pengembangan, terdapat beberapa masukan yang belum dapat diselesaikan dan akan menjadi catatan untuk pengembangan selanjutnya (fase berikutnya):')
p_backlog = document.add_paragraph('Fitur user dari mahasiswa saat ini belum terintegrasi dengan Program Studi (Prodi) yang menjadi sasaran aduan mahasiswa. Hal ini membutuhkan penyesuaian alur data agar setiap aduan mahasiswa dapat diteruskan secara presisi ke prodi masing-masing.')
p_backlog.style = 'List Number'

document.add_paragraph()
p = document.add_paragraph()
p.add_run('D. PENUTUP').bold = True
document.add_paragraph('Demikian laporan kegiatan dan pertanggungjawaban ini dibuat untuk dapat dipergunakan sebagaimana mestinya. Atas perhatian dan kerjasamanya, diucapkan terima kasih.')

document.add_paragraph()
document.add_paragraph()

# Signatures
table = document.add_table(rows=1, cols=2)
table.autofit = False

cell_left = table.cell(0, 0)
cell_right = table.cell(0, 1)

p_left = cell_left.paragraphs[0]
r_left = p_left.add_run('Mengetahui,\nDekan FTI UII\n\n\n\n\nProf. Dr. Sri Kusumadewi, S.Si., M.T.\nNIK. 945230102')
p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER

p_right = cell_right.paragraphs[0]
r_right = p_right.add_run('Yogyakarta, 7 Desember 2026\nPelaksana Tugas\n\n\n\n\nZulfahmi Nur Kesuma Atmaja, A.Md.')
p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER

document.save('Laporan_Kegiatan_Lapor_FTI.docx')
print("Document generated successfully.")
